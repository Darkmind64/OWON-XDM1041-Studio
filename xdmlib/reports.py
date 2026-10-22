# -*- coding: utf-8 -*-
"""Exports : rapport PDF d'intervention (Qt, sans dépendance) et classeur Excel avec graphique (openpyxl)."""
import datetime, html
from PySide6 import QtCore, QtGui


def build_pdf(path, info, sections, images):
    """info: dict titre/champs ; sections: liste (titre, html) ; images: dict nom -> QImage."""
    e = html.escape
    rows = "".join(f"<tr><td style='color:#555;padding-right:14px'>{e(k)}</td><td><b>{e(str(v))}</b></td></tr>"
                   for k, v in info["fields"] if v)
    body = [f"<h1 style='color:#1d6b3a'>{e(info['title'])}</h1>",
            f"<p style='color:#666'>Généré le {datetime.datetime.now():%d/%m/%Y à %H:%M} — {e(info.get('soft', ''))}</p>",
            f"<table cellspacing=2>{rows}</table>"]
    if info.get("notes"):
        body.append(f"<h3>Notes</h3><p>{e(info['notes']).replace(chr(10), '<br>')}</p>")
    for title, content in sections:
        body.append(f"<h3 style='color:#1d6b3a;border-bottom:1px solid #ccc'>{e(title)}</h3>{content}")
    doc = QtGui.QTextDocument()
    doc.setDefaultFont(QtGui.QFont("Segoe UI", 9))
    for name, img in images.items():
        doc.addResource(QtGui.QTextDocument.ResourceType.ImageResource, QtCore.QUrl(name), img)
    doc.setHtml("<html><body style='font-family:Segoe UI,Arial'>" + "".join(body) + "</body></html>")
    w = QtGui.QPdfWriter(path)
    w.setPageSize(QtGui.QPageSize(QtGui.QPageSize.PageSizeId.A4))
    w.setPageMargins(QtCore.QMarginsF(15, 15, 15, 15), QtGui.QPageLayout.Unit.Millimeter)
    w.setTitle(info["title"])
    w.setCreator(info.get("soft", ""))
    doc.print_(w)


def table_html(headers, rows, max_rows=200):
    e = html.escape
    h = "".join(f"<th style='background:#e8f3ec;padding:3px 6px;text-align:left'>{e(x)}</th>" for x in headers)
    b = "".join("<tr>" + "".join(f"<td style='padding:2px 6px;border-bottom:1px solid #eee'>{e(str(c))}</td>" for c in r) + "</tr>"
                for r in rows[:max_rows])
    more = f"<p style='color:#888'>… {len(rows) - max_rows} lignes supplémentaires non affichées</p>" if len(rows) > max_rows else ""
    return f"<table cellspacing=0 width='100%'><tr>{h}</tr>{b}</table>{more}"


def export_xlsx(path, times, values, unit, mode, stats, captures, cap_headers, events, ev_headers):
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Mesures"
    ws.append(["Horodatage", "Temps (s)", f"Valeur ({unit})"])
    t0 = times[0] if times else 0
    for t, v in zip(times, values):
        ws.append([datetime.datetime.fromtimestamp(t), round(t - t0, 3), v if v == v and abs(v) != float("inf") else None])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D6B3A")
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 16
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "dd/mm/yyyy hh:mm:ss.000"
    n = len(values)
    if n > 1:
        ch = LineChart()
        ch.title = f"{mode} ({unit})"
        ch.y_axis.title = unit
        ch.x_axis.title = "Temps (s)"
        ch.height, ch.width = 10, 24
        ch.add_data(Reference(ws, min_col=3, min_row=1, max_row=n + 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=2, min_row=2, max_row=n + 1))
        ch.series[0].smooth = False
        ch.series[0].graphicalProperties.line.width = 12000
        ch.legend = None
        ws.add_chart(ch, "E2")
    st = wb.create_sheet("Statistiques")
    for k, v in stats:
        st.append([k, v])
    st.column_dimensions["A"].width = 22
    st.column_dimensions["B"].width = 22
    for title, hdr, rows in (("Captures", cap_headers, captures), ("Événements", ev_headers, events)):
        if rows:
            s = wb.create_sheet(title)
            s.append(hdr)
            for r in rows:
                s.append(r)
            for c in s[1]:
                c.font = Font(bold=True)
    wb.save(path)
